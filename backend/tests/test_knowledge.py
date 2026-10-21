"""Knowledge base unit tests: extraction, chunking and the service with fakes."""

from __future__ import annotations

import io
import uuid
from contextlib import asynccontextmanager
from unittest.mock import AsyncMock

import pytest

from voice_agent.knowledge.chunk import chunk_text
from voice_agent.knowledge.extract import UnsupportedFileError, extract_text
from voice_agent.knowledge.service import EmptyDocumentError, KnowledgeService
from voice_agent.knowledge.store import Hit

TAMIL = "ரேஷன் கார்டு புதுப்பிக்க ஆதார் அட்டை தேவை. " * 5


def test_chunk_packs_paragraphs_and_overlaps():
    text = "\n\n".join(f"Paragraph {i}. " + "word " * 40 for i in range(10))
    chunks = chunk_text(text, max_chars=600, overlap=100)
    assert len(chunks) > 1
    assert all(len(c.text) <= 600 for c in chunks)
    assert [c.index for c in chunks] == list(range(len(chunks)))
    # the overlap repeats the tail of the previous chunk
    assert chunks[0].text[-60:].strip()[:20] in chunks[1].text


def test_chunk_splits_oversized_paragraph_by_sentence():
    text = "Sentence one is here. " * 100  # one paragraph, ~2200 chars
    chunks = chunk_text(text, max_chars=500, overlap=50)
    assert len(chunks) >= 5
    assert all(len(c.text) <= 500 for c in chunks)
    assert all(c.text.endswith(".") for c in chunks[:-1])


def test_chunk_handles_tamil_and_empty():
    assert chunk_text("   \n\n  ") == []
    [chunk] = chunk_text(TAMIL, max_chars=1200, overlap=100)
    assert "ரேஷன்" in chunk.text


def test_chunk_rejects_bad_sizes():
    with pytest.raises(ValueError):
        chunk_text("x", max_chars=100, overlap=100)


def test_extract_text_markdown_and_csv():
    assert extract_text("notes.md", b"# Title\n\nBody") == "# Title\n\nBody"
    csv_text = extract_text("fees.csv", b"service,fee\nrenewal,50\nnew card,100\n")
    assert csv_text == "service: renewal; fee: 50\nservice: new card; fee: 100"


def test_extract_text_docx():
    import docx

    document = docx.Document()
    document.add_paragraph("முதல் பத்தி")
    document.add_paragraph("Second paragraph")
    buf = io.BytesIO()
    document.save(buf)
    assert extract_text("doc.docx", buf.getvalue()) == "முதல் பத்தி\n\nSecond paragraph"


def test_extract_text_rejects_unknown_types():
    with pytest.raises(UnsupportedFileError, match="Unsupported"):
        extract_text("image.png", b"...")


class FakeEmbedder:
    model = "fake-embed"
    dimensions = 3

    def __init__(self):
        self.calls: list[list[str]] = []

    async def embed(self, texts):
        self.calls.append(texts)
        return [[float(len(t)), 1.0, 0.0] for t in texts]


class FakeStore:
    collection = "test"
    url = "http://fake-qdrant:6333"

    def __init__(self):
        self.points: dict[str, list] = {}
        self.names: dict[str, str] = {}
        self.ensure = AsyncMock()

    async def ensure_collection(self):
        await self.ensure()

    async def upsert(self, document_id, name, chunks, vectors):
        self.names[document_id] = name
        self.points[document_id] = list(zip(chunks, vectors, strict=True))

    async def search(self, vector, *, limit, document_id=None):
        hits = []
        for doc_id, items in self.points.items():
            if document_id and doc_id != document_id:
                continue
            for chunk, _ in items:
                hits.append(
                    Hit(0.9, chunk.text, doc_id, self.names.get(doc_id, "doc"), chunk.index)
                )
        return hits[:limit]

    async def delete_document(self, document_id):
        self.points.pop(document_id, None)


class FakePool:
    def __init__(self):
        self.statements: list[tuple[str, tuple]] = []

    @asynccontextmanager
    async def acquire(self):
        conn = AsyncMock()
        conn.execute.side_effect = lambda sql, *a: self.statements.append((sql, a))
        conn.fetchval.side_effect = lambda sql, *a: a[0] if "delete" in sql else None
        yield conn


async def test_service_ingests_and_searches():
    pool, store, embedder = FakePool(), FakeStore(), FakeEmbedder()
    service = KnowledgeService(
        pool=pool, store=store, embedder=embedder, chunk_chars=300, chunk_overlap=40
    )
    text = "\n\n".join(f"Fact number {i}: " + "detail " * 30 for i in range(6))

    result = await service.ingest("facts.txt", text.encode(), "text/plain")

    assert result.chunk_count == len(store.points[result.document_id]) > 1
    assert embedder.calls and len(embedder.calls[0]) == result.chunk_count
    statuses = [a[1] for sql, a in pool.statements if sql.startswith("update documents")]
    assert statuses == ["ready"]

    found = await service.search("fact", limit=3)
    assert len(found.hits) == 3 and found.hits[0].document_id == result.document_id
    assert await service.delete(result.document_id) is True
    assert result.document_id not in store.points


async def test_service_rejects_empty_documents_and_records_failures():
    pool, store, embedder = FakePool(), FakeStore(), FakeEmbedder()
    service = KnowledgeService(pool=pool, store=store, embedder=embedder)
    with pytest.raises(EmptyDocumentError):
        await service.ingest("empty.txt", b"   ", "text/plain")

    embedder.embed = AsyncMock(side_effect=RuntimeError("quota"))
    with pytest.raises(RuntimeError):
        await service.ingest("facts.txt", b"some facts here", "text/plain")
    failed = [
        a for sql, a in pool.statements if sql.startswith("update documents") and a[1] == "failed"
    ]
    assert failed and "quota" in failed[0][2]


def test_hit_is_json_serialisable():
    hit = Hit(0.5, "t", str(uuid.uuid4()), "d", 0)
    assert set(hit.__dict__) == {"score", "text", "document_id", "document_name", "chunk_index"}


def test_needs_translation_skips_english():
    from voice_agent.knowledge.translate import needs_translation

    assert needs_translation("ரேஷன் கார்டு புதுப்பிக்க என்ன தேவை?")
    assert not needs_translation("Is Aadhaar required for a new smart card?")
    assert not needs_translation("TNPDS 1967")  # digits only: nothing to translate


class FakeTranslator:
    def __init__(self):
        self.calls: list[str] = []

    async def to_english(self, text):
        self.calls.append(text)
        return "Is Aadhaar needed to renew a ration card?"


async def test_search_translates_indic_questions_before_embedding():
    store, embedder, translator = FakeStore(), FakeEmbedder(), FakeTranslator()
    service = KnowledgeService(pool=None, store=store, embedder=embedder, translator=translator)
    await service.search("ரேஷன் கார்டை புதுப்பிக்க ஆதார் தேவையா?", limit=2)
    assert translator.calls == ["ரேஷன் கார்டை புதுப்பிக்க ஆதார் தேவையா?"]
    assert embedder.calls[-1] == ["Is Aadhaar needed to renew a ration card?"]

    await service.search("Is Aadhaar needed?", limit=2)
    assert len(translator.calls) == 1  # English is embedded as-is
    assert embedder.calls[-1] == ["Is Aadhaar needed?"]


async def test_knowledge_tool_handler_returns_passages_and_speaks_filler():
    from unittest.mock import Mock

    from voice_agent.knowledge.tool import KnowledgeTool, searching_line

    store, embedder = FakeStore(), FakeEmbedder()
    service = KnowledgeService(pool=None, store=store, embedder=embedder)
    from voice_agent.knowledge.chunk import Chunk

    await store.upsert(
        "doc-1", "faq.pdf", [Chunk(0, "Aadhaar is mandatory for a new smart card.")], [[1.0, 0, 0]]
    )
    tool = KnowledgeTool(service, limit=3)
    results = []
    params = Mock(
        arguments={"question": "ஆதார் தேவையா?"},
        result_callback=AsyncMock(side_effect=lambda r: results.append(r)),
    )
    await tool.handle(params)
    assert results[0]["found"] is True
    assert results[0]["passages"][0]["source"] == "faq.pdf"
    assert "Aadhaar" in results[0]["passages"][0]["text"]

    empty = Mock(arguments={}, result_callback=AsyncMock(side_effect=lambda r: results.append(r)))
    await tool.handle(empty)
    assert results[-1] == {"found": False, "reason": "empty question"}

    assert "நிமிடம்" in searching_line("ta-IN")
    assert searching_line("xx-XX") == searching_line("en-IN")
    assert tool.schema().standard_tools[0].name == "search_knowledge_base"


async def test_knowledge_tool_notifies_listeners_with_chunks():
    from unittest.mock import Mock

    from voice_agent.knowledge.chunk import Chunk
    from voice_agent.knowledge.tool import KnowledgeTool

    store, embedder = FakeStore(), FakeEmbedder()
    await store.upsert("doc-1", "faq.pdf", [Chunk(0, "Aadhaar is mandatory.")], [[1.0, 0, 0]])
    tool = KnowledgeTool(KnowledgeService(pool=None, store=store, embedder=embedder))
    records = []
    tool.on_search(records.append)
    params = Mock(arguments={"question": "Is Aadhaar needed?"}, result_callback=AsyncMock())
    await tool.handle(params)
    [record] = records
    assert record.question == "Is Aadhaar needed?"
    assert record.hits[0]["source"] == "faq.pdf" and record.hits[0]["used"] is True
    assert set(record.as_dict()) >= {"question", "translated_query", "hits", "embed_ms", "at"}


async def test_tool_failure_names_the_qdrant_url():
    from unittest.mock import Mock

    from voice_agent.knowledge.tool import KnowledgeTool

    class BrokenStore(FakeStore):
        url = "http://qdrant:6333"

        async def search(self, *a, **k):
            raise OSError("[Errno -3] Temporary failure in name resolution")

    service = KnowledgeService(pool=None, store=BrokenStore(), embedder=FakeEmbedder())
    tool = KnowledgeTool(service)
    records, results = [], []
    tool.on_search(records.append)
    await tool.handle(
        Mock(
            arguments={"question": "q"},
            result_callback=AsyncMock(side_effect=lambda r: results.append(r)),
        )
    )
    assert "qdrant=http://qdrant:6333" in records[0].error
    assert results[0] == {"found": False, "reason": "search unavailable"}


def test_whole_document_mode_keeps_one_chunk_until_too_long():
    from voice_agent.knowledge.chunk import WHOLE_DOCUMENT_MAX_CHARS

    text = "\n\n".join(f"Step {i}: " + "detail " * 50 for i in range(8))  # ~2.5k chars
    [chunk] = chunk_text(text, whole_document=True)
    assert chunk.text.startswith("Step 0") and "Step 7" in chunk.text
    huge = "para " * (WHOLE_DOCUMENT_MAX_CHARS // 4)
    assert len(chunk_text(huge, max_chars=1200, overlap=200, whole_document=True)) > 1
    assert chunk_text("   ", whole_document=True) == []


async def test_service_whole_document_ingest_is_one_chunk():
    pool, store, embedder = FakePool(), FakeStore(), FakeEmbedder()
    service = KnowledgeService(pool=pool, store=store, embedder=embedder, whole_document=True)
    text = "\n\n".join(
        f"Fact {i}: " + "detail " * 60 for i in range(6)
    )  # would be 3+ chunks otherwise
    result = await service.ingest("sheet.txt", text.encode(), "text/plain")
    assert result.chunk_count == 1
